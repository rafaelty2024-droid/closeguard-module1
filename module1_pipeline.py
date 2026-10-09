from collections import Counter
from datetime import datetime, timezone
from email import message_from_bytes
from email.header import decode_header
from pathlib import Path
import hashlib
import imaplib
import json
import os
import re

import numpy as np
import pandas as pd
import pdfplumber
from dotenv import load_dotenv

try:
    from PIL import Image, ImageOps, ImageSequence
except ImportError:  # The dashboard will explain the missing OCR dependency.
    Image = None
    ImageOps = None
    ImageSequence = None

try:
    import pytesseract
except ImportError:
    pytesseract = None

try:
    from rapidocr import RapidOCR
except ImportError:
    RapidOCR = None

try:
    import pymupdf
except ImportError:
    pymupdf = None

try:
    from docx import Document
except ImportError:
    Document = None

try:
    from pillow_heif import register_heif_opener
except ImportError:
    register_heif_opener = None


BASE_FOLDER = Path(__file__).resolve().parent
load_dotenv(BASE_FOLDER / ".env", override=True)

INPUT_FOLDER = BASE_FOLDER / "input"
OUTPUT_FOLDER = BASE_FOLDER / "output"
REVIEW_FOLDER = BASE_FOLDER / "review"
REFERENCE_FOLDER = BASE_FOLDER / "reference_samples"

HASH_FILE = BASE_FOLDER / "processed_hashes.txt"
ATTACHMENT_REGISTRY_FILE = OUTPUT_FOLDER / "attachment_registry.csv"
CONTENT_REGISTRY_FILE = OUTPUT_FOLDER / "content_fingerprints.csv"

INTAKE_FILE = OUTPUT_FOLDER / "intake_register.csv"
MODULE2_FILE = OUTPUT_FOLDER / "module2_input.csv"
MODULE2_JSON_FILE = OUTPUT_FOLDER / "module2_input.json"
REVIEW_FILE = REVIEW_FOLDER / "intake_review.csv"
UNSUPPORTED_FILE = REVIEW_FOLDER / "unsupported_documents.csv"
DUPLICATE_FILE = REVIEW_FOLDER / "duplicate_attachments.csv"

IMAGE_EXTENSIONS = {
    ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp",
    ".heic", ".heif",
}

SUPPORTED_EXTENSIONS = {
    ".pdf", ".csv", ".xlsx", ".txt", ".docx", *IMAGE_EXTENSIONS,
}

# JSON is intentionally reference-only. An emailed JSON attachment is not
# treated as transaction evidence or passed to Module 2.
REFERENCE_EXTENSIONS = {
    ".pdf", ".txt", ".docx", ".json", *IMAGE_EXTENSIONS,
}

MAX_FILE_SIZE = 20 * 1024 * 1024
MIN_USABLE_TEXT_CHARACTERS = 25

STANDARD_COLUMNS = [
    "record_id",
    "document_type",
    "supplier",
    "document_number",
    "po_reference",
    "document_date",
    "amount",
    "currency",
    "source_file",
    "source_extension",
    "source_hash",
    "content_fingerprint",
    "source_row",
    "email_id",
    "extraction_method",
    "ocr_used",
    "classification_method",
    "matched_reference_sample",
    "processing_status",
    "missing_fields",
    "intake_note",
]

DUPLICATE_COLUMNS = [
    "detected_at",
    "email_id",
    "source_filename",
    "source_hash",
    "duplicate_of",
    "duplicate_reason",
]

DEFAULT_REFERENCE_CONFIG = {
    "classification_threshold": 0.18,
    "document_type_aliases": {
        "PURCHASE_ORDER": [
            "purchase order", "purchase order document", "po",
        ],
        "SUPPLIER_INVOICE": [
            "supplier invoice", "sales invoice", "tax invoice",
            "commercial invoice", "invoice",
        ],
        "RECEIPT": ["receipt", "payment receipt"],
    },
    "field_labels": {
        "document_type": ["Document Type", "Document Category", "Type"],
        "supplier": [
            "Supplier", "Supplier Name", "Vendor", "Vendor Name",
            "Issued By", "Seller", "Sold By", "From",
        ],
        "invoice_number": [
            "Document Number", "Invoice Number", "Invoice No", "Invoice #",
            "Sales Invoice Number", "Tax Invoice Number", "Bill Number",
        ],
        "purchase_order_number": [
            "Document Number", "PO Number", "PO No", "PO #",
            "Purchase Order Number", "Purchase Order No", "Order Number",
        ],
        "po_reference": [
            "PO Reference", "Purchase Order Reference", "PO Number", "PO No",
            "PO #", "Purchase Order Number", "Customer PO", "Buyer PO",
        ],
        "document_date": [
            "Document Date", "Invoice Date", "Order Date", "PO Date", "Date",
        ],
        "amount": [
            "Amount", "Total Amount", "Invoice Total", "Order Total",
            "Grand Total", "Amount Due", "Net Total", "Total",
        ],
        "currency": ["Currency", "Currency Code"],
    },
}

REFERENCE_STOP_WORDS = {
    "THE", "AND", "FOR", "WITH", "FROM", "THIS", "THAT", "YOUR", "OUR",
    "PAGE", "PLEASE", "TOTAL", "DATE", "DOCUMENT", "AMOUNT", "NUMBER",
}

_RAPID_OCR_ENGINE = None
_REFERENCE_PROFILE_CACHE = {
    "signature": None,
    "profiles": [],
}


def utc_now():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def create_folders():
    for folder in (
        INPUT_FOLDER,
        OUTPUT_FOLDER,
        REVIEW_FOLDER,
        REFERENCE_FOLDER / "purchase_orders",
        REFERENCE_FOLDER / "supplier_invoices",
    ):
        folder.mkdir(parents=True, exist_ok=True)


def decode_email_text(value):
    if not value:
        return ""
    output = []
    for part, encoding in decode_header(value):
        if isinstance(part, bytes):
            output.append(part.decode(encoding or "utf-8", errors="replace"))
        else:
            output.append(part)
    return "".join(output)


def safe_filename(filename):
    filename = Path(filename).name
    return re.sub(r"[^A-Za-z0-9._-]", "_", filename)


def calculate_hash(content):
    return hashlib.sha256(content).hexdigest()


def create_record_id(source_hash, source_row=None, document_number=None):
    identity = f"{source_hash}|{source_row or 1}|{document_number or ''}"
    return hashlib.sha256(identity.encode("utf-8")).hexdigest()[:16]


def load_processed_hashes():
    if not HASH_FILE.exists():
        return set()
    return {
        line.strip()
        for line in HASH_FILE.read_text(encoding="utf-8").splitlines()
        if line.strip()
    }


def save_processed_hash(hash_value):
    with HASH_FILE.open("a", encoding="utf-8") as file:
        file.write(hash_value + "\n")


def load_csv_safely(path, columns):
    if not path.exists():
        return pd.DataFrame(columns=columns)
    try:
        data = pd.read_csv(path, dtype=str).fillna("")
    except (pd.errors.EmptyDataError, pd.errors.ParserError):
        return pd.DataFrame(columns=columns)
    for column in columns:
        if column not in data.columns:
            data[column] = ""
    return data[columns]


def append_registry_row(path, row, columns):
    data = load_csv_safely(path, columns)
    data = pd.concat([data, pd.DataFrame([row], columns=columns)], ignore_index=True)
    data.to_csv(path, index=False)


def load_attachment_registry():
    data = load_csv_safely(
        ATTACHMENT_REGISTRY_FILE,
        ["source_hash", "stored_filename", "first_email_id", "original_filename", "received_at"],
    )
    return {
        row["source_hash"]: row.to_dict()
        for _, row in data.iterrows()
        if row["source_hash"]
    }


def register_attachment(source_hash, stored_filename, email_id, original_filename):
    append_registry_row(
        ATTACHMENT_REGISTRY_FILE,
        {
            "source_hash": source_hash,
            "stored_filename": stored_filename,
            "first_email_id": email_id,
            "original_filename": original_filename,
            "received_at": utc_now(),
        },
        ["source_hash", "stored_filename", "first_email_id", "original_filename", "received_at"],
    )


def log_duplicate(email_id, filename, source_hash, duplicate_of, reason):
    append_registry_row(
        DUPLICATE_FILE,
        {
            "detected_at": utc_now(),
            "email_id": email_id,
            "source_filename": filename,
            "source_hash": source_hash,
            "duplicate_of": duplicate_of,
            "duplicate_reason": reason,
        },
        DUPLICATE_COLUMNS,
    )


def load_content_registry():
    data = load_csv_safely(
        CONTENT_REGISTRY_FILE,
        ["content_fingerprint", "source_hash", "source_file", "email_id", "registered_at"],
    )
    return {
        row["content_fingerprint"]: row.to_dict()
        for _, row in data.iterrows()
        if row["content_fingerprint"]
    }


def register_content_fingerprint(fingerprint, source_hash, source_file, email_id):
    append_registry_row(
        CONTENT_REGISTRY_FILE,
        {
            "content_fingerprint": fingerprint,
            "source_hash": source_hash,
            "source_file": source_file,
            "email_id": email_id,
            "registered_at": utc_now(),
        },
        ["content_fingerprint", "source_hash", "source_file", "email_id", "registered_at"],
    )


def calculate_content_fingerprint(text):
    canonical = re.sub(r"[^A-Z0-9]+", "", str(text).upper())
    if len(canonical) < MIN_USABLE_TEXT_CHARACTERS:
        return None
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()


def download_email_attachments():
    """Download unique supported attachments from unread matching emails."""
    create_folders()
    email_host = os.getenv("EMAIL_HOST")
    email_address = os.getenv("EMAIL_ADDRESS")
    email_password = os.getenv("EMAIL_PASSWORD")
    subject_filter = os.getenv("EMAIL_SUBJECT_FILTER", "").strip()

    if not email_host:
        raise ValueError("EMAIL_HOST is missing from the .env file.")
    if not email_address:
        raise ValueError("EMAIL_ADDRESS is missing from the .env file.")
    if not email_password:
        raise ValueError("EMAIL_PASSWORD is missing from the .env file.")

    processed_hashes = load_processed_hashes()
    attachment_registry = load_attachment_registry()
    downloaded_files = []
    mailbox = None

    try:
        mailbox = imaplib.IMAP4_SSL(email_host)
        mailbox.login(email_address, email_password)
        mailbox.select("INBOX")
        status, message_numbers = mailbox.search(None, "UNSEEN")
        if status != "OK":
            raise RuntimeError("The inbox could not be searched.")

        for message_number in message_numbers[0].split():
            status, message_data = mailbox.fetch(message_number, "(RFC822)")
            if status != "OK":
                continue
            message = message_from_bytes(message_data[0][1])
            subject = decode_email_text(message.get("Subject"))
            if subject_filter and subject_filter.upper() not in subject.upper():
                continue

            email_id = message_number.decode()
            handled_supported_attachment = False

            for part in message.walk():
                original_filename = part.get_filename()
                if not original_filename:
                    continue
                filename = safe_filename(decode_email_text(original_filename))
                extension = Path(filename).suffix.lower()
                if extension not in SUPPORTED_EXTENSIONS:
                    continue
                content = part.get_payload(decode=True)
                if not content or len(content) > MAX_FILE_SIZE:
                    continue

                handled_supported_attachment = True
                source_hash = calculate_hash(content)
                if source_hash in processed_hashes:
                    previous = attachment_registry.get(source_hash, {})
                    log_duplicate(
                        email_id=email_id,
                        filename=filename,
                        source_hash=source_hash,
                        duplicate_of=previous.get("stored_filename", "previously processed attachment"),
                        reason="EXACT_FILE_HASH",
                    )
                    continue

                output_name = f"{email_id}_{source_hash[:10]}_{filename}"
                output_path = INPUT_FOLDER / output_name
                output_path.write_bytes(content)
                downloaded_files.append(
                    {"path": output_path, "email_id": email_id, "source_hash": source_hash}
                )
                save_processed_hash(source_hash)
                processed_hashes.add(source_hash)
                register_attachment(source_hash, output_name, email_id, filename)
                attachment_registry[source_hash] = {"stored_filename": output_name}

            if handled_supported_attachment:
                mailbox.store(message_number, "+FLAGS", "\\Seen")
    finally:
        if mailbox is not None:
            try:
                mailbox.logout()
            except Exception:
                pass
    return downloaded_files


def is_usable_text(text):
    return len(re.sub(r"\W", "", text or "")) >= MIN_USABLE_TEXT_CHARACTERS


def tesseract_available():
    if pytesseract is None:
        return False
    try:
        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False


def ocr_engine_name():
    if Image is None:
        return None
    if RapidOCR is not None:
        return "RAPIDOCR_ONNX"
    if tesseract_available():
        return "TESSERACT"
    return None


def ocr_available():
    return ocr_engine_name() is not None


def get_rapid_ocr_engine():
    global _RAPID_OCR_ENGINE
    if RapidOCR is None:
        return None
    if _RAPID_OCR_ENGINE is None:
        _RAPID_OCR_ENGINE = RapidOCR()
    return _RAPID_OCR_ENGINE


def prepare_image_for_ocr(image):
    image = image.convert("RGB")
    grayscale = ImageOps.grayscale(image)
    grayscale = ImageOps.autocontrast(grayscale)
    if grayscale.width < 1800:
        scale = 1800 / grayscale.width
        grayscale = grayscale.resize(
            (int(grayscale.width * scale), int(grayscale.height * scale))
        )
    return grayscale


def ocr_image(image):
    if not ocr_available():
        raise RuntimeError(
            "OCR dependencies are unavailable. Install the project requirements."
        )
    prepared = prepare_image_for_ocr(image)
    if RapidOCR is not None:
        engine = get_rapid_ocr_engine()
        result = engine(np.asarray(prepared.convert("RGB")))
        lines = getattr(result, "txts", None) or ()
        return "\n".join(str(line) for line in lines if str(line).strip()).strip()
    return pytesseract.image_to_string(prepared, config="--psm 6").strip()


def extract_pdf_text(file_path):
    pages = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text() or "")
    direct_text = "\n".join(pages).strip()
    if is_usable_text(direct_text):
        return direct_text, "PDF_TEXT", False

    if pymupdf is None or not ocr_available():
        return direct_text, "PDF_TEXT_NO_OCR", False

    ocr_pages = []
    with pymupdf.open(file_path) as document:
        for page in document:
            pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2.2, 2.2), alpha=False)
            image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
            ocr_pages.append(ocr_image(image))
    return "\n".join(ocr_pages).strip(), "PDF_OCR", True


def extract_image_text(file_path):
    if Image is None:
        return "", "IMAGE_NO_OCR", False
    if file_path.suffix.lower() in {".heic", ".heif"} and register_heif_opener:
        register_heif_opener()
    if not ocr_available():
        return "", "IMAGE_NO_OCR", False
    pages = []
    with Image.open(file_path) as image:
        for frame in ImageSequence.Iterator(image):
            pages.append(ocr_image(frame.copy()))
    return "\n".join(pages).strip(), "IMAGE_OCR", True


def extract_docx_text(file_path):
    if Document is None:
        raise RuntimeError("python-docx is required to read DOCX files.")
    document = Document(file_path)
    output = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            output.append("  ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(output).strip(), "DOCX_TEXT", False


def extract_text_from_path(file_path):
    extension = file_path.suffix.lower()
    if extension == ".pdf":
        return extract_pdf_text(file_path)
    if extension in IMAGE_EXTENSIONS:
        return extract_image_text(file_path)
    if extension == ".docx":
        return extract_docx_text(file_path)
    if extension == ".txt":
        return file_path.read_text(encoding="utf-8", errors="replace").strip(), "TEXT_FILE", False
    raise ValueError(f"Text extraction is not supported for {extension}.")


def merge_config(base, override):
    output = dict(base)
    for key, value in override.items():
        if isinstance(value, dict) and isinstance(output.get(key), dict):
            output[key] = merge_config(output[key], value)
        else:
            output[key] = value
    return output


def normalise_reference_document_type(value):
    """Convert JSON reference type names into Module 1 categories."""
    if value is None:
        return None
    normalised = re.sub(r"[^A-Z0-9]+", "_", str(value).upper()).strip("_")
    if normalised in {"PO", "PURCHASE_ORDER", "PURCHASE_ORDER_DOCUMENT"}:
        return "PURCHASE_ORDER"
    if normalised in {
        "INVOICE",
        "SALES_INVOICE",
        "SUPPLIER_INVOICE",
        "TAX_INVOICE",
        "COMMERCIAL_INVOICE",
    }:
        return "SUPPLIER_INVOICE"
    return normalised


def read_json_reference(file_path, expected_document_type=None):
    """Read and validate one structured JSON reference sample."""
    try:
        data = json.loads(file_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None

    if not isinstance(data, dict):
        return None

    document_type = normalise_reference_document_type(data.get("document_type"))
    if expected_document_type and document_type and document_type != expected_document_type:
        return None

    field_values = data.get("field_values", {})
    labels = data.get("labels", {})
    keywords = data.get("keywords", [])

    if not isinstance(field_values, dict) or not isinstance(labels, dict):
        return None
    if isinstance(keywords, str):
        keywords = [keywords]
    if not isinstance(keywords, list):
        return None

    clean_labels = {}
    for field, aliases in labels.items():
        if isinstance(aliases, str):
            aliases = [aliases]
        if not isinstance(aliases, list):
            return None
        clean_labels[field] = [
            str(alias).strip()
            for alias in aliases
            if str(alias).strip()
        ]

    return {
        "document_type": document_type or expected_document_type,
        "field_values": field_values,
        "labels": clean_labels,
        "keywords": [str(keyword).strip() for keyword in keywords if str(keyword).strip()],
    }


def json_reference_text(reference_data):
    """Create deterministic comparison text from structured reference data."""
    output = []
    document_type = reference_data.get("document_type")
    if document_type:
        output.append(f"Document Type {document_type}")

    for field, value in sorted(reference_data.get("field_values", {}).items()):
        if value is not None and str(value).strip():
            output.append(f"{field.replace('_', ' ')} {value}")

    for field, aliases in sorted(reference_data.get("labels", {}).items()):
        for alias in aliases:
            output.append(f"{field.replace('_', ' ')} {alias}")

    output.extend(reference_data.get("keywords", []))
    return "\n".join(output)


def iter_reference_files():
    folders = {
        "PURCHASE_ORDER": REFERENCE_FOLDER / "purchase_orders",
        "SUPPLIER_INVOICE": REFERENCE_FOLDER / "supplier_invoices",
    }
    for document_type, folder in folders.items():
        if not folder.exists():
            continue
        for path in sorted(folder.iterdir()):
            if path.is_file() and path.suffix.lower() in REFERENCE_EXTENSIONS:
                yield document_type, path


def merge_json_reference_labels(config):
    """Add approved supplier-specific JSON labels to extraction aliases."""
    allowed_fields = config.get("field_labels", {})
    for document_type, path in iter_reference_files():
        if path.suffix.lower() != ".json":
            continue
        reference_data = read_json_reference(path, document_type)
        if not reference_data:
            continue
        for field, aliases in reference_data["labels"].items():
            if field not in allowed_fields:
                continue
            current = allowed_fields[field]
            existing = {str(value).casefold() for value in current}
            for alias in aliases:
                if alias.casefold() not in existing:
                    current.append(alias)
                    existing.add(alias.casefold())
    return config


def load_reference_config():
    # Copy the defaults so dynamically merged JSON labels never mutate the
    # module-level configuration shared by later Streamlit reruns.
    config = json.loads(json.dumps(DEFAULT_REFERENCE_CONFIG))
    path = REFERENCE_FOLDER / "reference_labels.json"
    if path.exists():
        try:
            config = merge_config(config, json.loads(path.read_text(encoding="utf-8")))
        except (json.JSONDecodeError, OSError):
            pass
    return merge_json_reference_labels(config)


def reference_tokens(text):
    counts = Counter(
        token
        for token in re.findall(r"[A-Z]{3,}", str(text).upper())
        if token not in REFERENCE_STOP_WORDS
    )
    return set(counts)


def build_reference_profiles():
    files = list(iter_reference_files())
    signature = tuple(
        (str(path), path.stat().st_mtime_ns, path.stat().st_size)
        for _, path in files
    )
    if _REFERENCE_PROFILE_CACHE["signature"] == signature:
        return _REFERENCE_PROFILE_CACHE["profiles"]

    profiles = []
    for document_type, path in files:
        try:
            if path.suffix.lower() == ".json":
                reference_data = read_json_reference(path, document_type)
                if not reference_data:
                    continue
                text = json_reference_text(reference_data)
            else:
                text, _, _ = extract_text_from_path(path)
        except Exception:
            continue
        tokens = reference_tokens(text)
        if tokens:
            profiles.append(
                {"document_type": document_type, "path": path, "tokens": tokens}
            )

    _REFERENCE_PROFILE_CACHE["signature"] = signature
    _REFERENCE_PROFILE_CACHE["profiles"] = profiles
    return profiles


def get_reference_summary():
    create_folders()
    reference_files = list(iter_reference_files())
    json_files = [
        (document_type, path)
        for document_type, path in reference_files
        if path.suffix.lower() == ".json"
    ]
    valid_json = sum(
        read_json_reference(path, document_type) is not None
        for document_type, path in json_files
    )
    return {
        "purchase_orders": len([
            path for document_type, path in reference_files
            if document_type == "PURCHASE_ORDER"
        ]),
        "supplier_invoices": len([
            path for document_type, path in reference_files
            if document_type == "SUPPLIER_INVOICE"
        ]),
        "json_samples": len(json_files),
        "invalid_json_samples": len(json_files) - valid_json,
        "ocr_available": ocr_available(),
        "ocr_engine": ocr_engine_name(),
        "supported_extensions": sorted(SUPPORTED_EXTENSIONS),
        "reference_extensions": sorted(REFERENCE_EXTENSIONS),
    }


def extract_labeled_value(labels, text):
    lines = [line.strip() for line in str(text).splitlines() if line.strip()]
    labels = sorted({str(label).strip() for label in labels if str(label).strip()}, key=len, reverse=True)
    for label in labels:
        escaped = re.escape(label)
        strict = re.compile(rf"^{escaped}\s*(?:[:#-]\s*|\s{{2,}})(.+)$", re.IGNORECASE)
        for line in lines:
            match = strict.match(line)
            if match and match.group(1).strip():
                return match.group(1).strip()
        prefix = label + " "
        for line in lines:
            if line.startswith(prefix):
                value = line[len(prefix):].strip()
                if value:
                    return value
        # OCR sometimes removes the single space between a printed label
        # and its value. Require the configured label's exact capitalisation
        # here so ordinary narrative text is not mistaken for a field.
        for line in lines:
            if line.startswith(label) and len(line) > len(label):
                next_character = line[len(label)]
                if next_character.isalnum() and not next_character.islower():
                    value = line[len(label):].strip()
                    if value:
                        return value
    return None


def normalise_type_from_value(value, aliases):
    if not value:
        return "UNKNOWN"
    normalised = re.sub(r"[^A-Z0-9]+", " ", str(value).upper()).strip()
    for document_type, values in aliases.items():
        for alias in values:
            alias_normalised = re.sub(r"[^A-Z0-9]+", " ", alias.upper()).strip()
            if normalised == alias_normalised or alias_normalised in normalised:
                return document_type
    return "UNKNOWN"


def classify_document(text, config, reference_profiles):
    labels = config["field_labels"]["document_type"]
    labelled_type = extract_labeled_value(labels, text)
    document_type = normalise_type_from_value(labelled_type, config["document_type_aliases"])
    if document_type != "UNKNOWN":
        return document_type, "EXPLICIT_DOCUMENT_TYPE", None

    lines = [line.strip() for line in str(text).splitlines() if line.strip()][:25]
    for line in lines:
        if re.match(r"^(?:TAX|SALES|SUPPLIER|COMMERCIAL)?\s*INVOICE(?:\b|\s*[-#])", line, re.IGNORECASE):
            return "SUPPLIER_INVOICE", "TITLE_PATTERN", None
        if re.match(r"^PURCHASE\s+ORDER(?:\b|\s*[-#])", line, re.IGNORECASE):
            return "PURCHASE_ORDER", "TITLE_PATTERN", None
        if re.match(r"^(?:PAYMENT\s+)?RECEIPT(?:\b|\s*[-#])", line, re.IGNORECASE):
            return "RECEIPT", "TITLE_PATTERN", None

    incoming = reference_tokens(text)
    label_groups = config["field_labels"]
    labelled_field_count = sum(
        extract_labeled_value(label_groups[field], text) is not None
        for field in (
            "supplier",
            "invoice_number",
            "purchase_order_number",
            "po_reference",
            "document_date",
            "amount",
        )
    )
    best_profile = None
    best_score = 0.0
    best_shared_count = 0
    for profile in reference_profiles:
        union = incoming | profile["tokens"]
        shared_count = len(incoming & profile["tokens"])
        score = shared_count / len(union) if union else 0.0
        if score > best_score:
            best_score = score
            best_shared_count = shared_count
            best_profile = profile
    if (
        best_profile
        and labelled_field_count >= 2
        and best_shared_count >= 3
        and best_score >= float(config.get("classification_threshold", 0.18))
    ):
        return (
            best_profile["document_type"],
            f"REFERENCE_SAMPLE_SIMILARITY_{best_score:.2f}",
            best_profile["path"].name,
        )
    return "UNKNOWN", "NO_RELIABLE_CLASSIFICATION", None


def parse_amount_and_currency(value):
    if not value:
        return None, None
    text = str(value).strip()
    currency_match = re.search(r"\b(SGD|USD|EUR|GBP|MYR|AUD|JPY|CNY)\b", text, re.IGNORECASE)
    currency = currency_match.group(1).upper() if currency_match else None
    if currency is None:
        if "S$" in text:
            currency = "SGD"
        elif "$" in text:
            currency = "USD"
        elif "€" in text:
            currency = "EUR"
        elif "£" in text:
            currency = "GBP"
    number_match = re.search(r"\(?-?\d[\d,]*(?:\.\d{1,2})?\)?", text)
    if not number_match:
        return None, currency
    number = number_match.group(0)
    negative = number.startswith("(") and number.endswith(")")
    number = number.strip("()").replace(",", "")
    try:
        amount = float(number)
        if negative:
            amount = -amount
        return round(amount, 2), currency
    except ValueError:
        return None, currency


def extract_fields_from_text(file_path, email_id, source_hash, text, extraction_method, ocr_used, config, profiles):
    document_type, classification_method, matched_sample = classify_document(text, config, profiles)
    labels = config["field_labels"]
    supplier = extract_labeled_value(labels["supplier"], text)
    if document_type == "PURCHASE_ORDER":
        number_labels = labels["purchase_order_number"]
    else:
        number_labels = labels["invoice_number"]
    document_number = extract_labeled_value(number_labels, text)
    po_reference = extract_labeled_value(labels["po_reference"], text)
    if document_type == "PURCHASE_ORDER" and not po_reference:
        po_reference = document_number
    document_date = extract_labeled_value(labels["document_date"], text)
    amount_value = extract_labeled_value(labels["amount"], text)
    amount, currency = parse_amount_and_currency(amount_value)
    explicit_currency = extract_labeled_value(labels["currency"], text)
    if explicit_currency:
        currency = explicit_currency

    return {
        "document_type": document_type,
        "supplier": supplier,
        "document_number": document_number,
        "po_reference": po_reference,
        "document_date": document_date,
        "amount": amount,
        "currency": currency,
        "source_file": file_path.name,
        "source_extension": file_path.suffix.lower(),
        "source_hash": source_hash,
        "source_row": None,
        "email_id": email_id,
        "text_extracted": is_usable_text(text),
        "extraction_method": extraction_method,
        "ocr_used": ocr_used,
        "classification_method": classification_method,
        "matched_reference_sample": matched_sample,
    }


def normalise_table_columns(dataframe):
    aliases = {
        "type": "document_type", "document type": "document_type", "document_type": "document_type",
        "vendor": "supplier", "vendor name": "supplier", "vendor_name": "supplier",
        "supplier name": "supplier", "supplier_name": "supplier", "supplier": "supplier",
        "document number": "document_number", "document_number": "document_number",
        "invoice number": "document_number", "invoice_number": "document_number",
        "invoice no": "document_number", "invoice_no": "document_number",
        "po reference": "po_reference", "po_reference": "po_reference",
        "po number": "po_reference", "po_number": "po_reference", "po no": "po_reference", "po_no": "po_reference",
        "document date": "document_date", "document_date": "document_date",
        "invoice date": "document_date", "invoice_date": "document_date", "date": "document_date",
        "total": "amount", "total amount": "amount", "total_amount": "amount", "amount": "amount",
        "currency": "currency",
    }
    renamed = {}
    for column in dataframe.columns:
        name = str(column).strip().lower()
        renamed[column] = aliases.get(name, name.replace(" ", "_"))
    return dataframe.rename(columns=renamed)


def read_table(file_path, email_id, source_hash):
    dataframe = pd.read_csv(file_path) if file_path.suffix.lower() == ".csv" else pd.read_excel(file_path)
    dataframe = normalise_table_columns(dataframe)
    records = []
    for index, row in dataframe.iterrows():
        record = row.to_dict()
        record.update(
            {
                "source_file": file_path.name,
                "source_extension": file_path.suffix.lower(),
                "source_hash": source_hash,
                "source_row": index + 2,
                "email_id": email_id,
                "text_extracted": True,
                "extraction_method": "CSV_TABLE" if file_path.suffix.lower() == ".csv" else "EXCEL_TABLE",
                "ocr_used": False,
                "classification_method": "TABLE_FIELD",
                "matched_reference_sample": None,
            }
        )
        records.append(record)
    canonical = dataframe.fillna("").astype(str).to_csv(index=False)
    return records, canonical


def standardise_text(value):
    if pd.isna(value) or not str(value).strip():
        return None
    value = re.sub(r"[^\w\s]", "", str(value).upper().strip())
    return re.sub(r"\s+", " ", value)


def standardise_reference(value):
    if pd.isna(value) or not str(value).strip():
        return None
    return re.sub(r"[^A-Z0-9]", "", str(value).upper())


def standardise_date(value):
    if pd.isna(value) or not str(value).strip():
        return None
    parsed = pd.to_datetime(value, errors="coerce", dayfirst=True)
    return None if pd.isna(parsed) else parsed.strftime("%Y-%m-%d")


def standardise_amount(value):
    if pd.isna(value) or not str(value).strip():
        return None
    if isinstance(value, (int, float)):
        return round(float(value), 2)
    amount, _ = parse_amount_and_currency(value)
    return amount


def standardise_currency(value):
    if pd.isna(value) or not str(value).strip():
        return None
    return str(value).upper().strip()


def normalise_document_type(value):
    if pd.isna(value) or not str(value).strip():
        return "UNKNOWN"
    return normalise_type_from_value(value, DEFAULT_REFERENCE_CONFIG["document_type_aliases"])


def clean_record(record, content_fingerprint=None):
    document_type = normalise_document_type(record.get("document_type"))
    supplier = standardise_text(record.get("supplier"))
    document_number = standardise_reference(record.get("document_number"))
    po_reference = standardise_reference(record.get("po_reference"))
    document_date = standardise_date(record.get("document_date"))
    amount = standardise_amount(record.get("amount"))
    currency = standardise_currency(record.get("currency"))
    source_hash = record.get("source_hash")
    source_row = record.get("source_row")
    record_id = create_record_id(source_hash, source_row, document_number)

    cleaned = {
        "record_id": record_id,
        "document_type": document_type,
        "supplier": supplier,
        "document_number": document_number,
        "po_reference": po_reference,
        "document_date": document_date,
        "amount": amount,
        "currency": currency,
        "source_file": record.get("source_file"),
        "source_extension": record.get("source_extension"),
        "source_hash": source_hash,
        "content_fingerprint": content_fingerprint,
        "source_row": source_row,
        "email_id": record.get("email_id"),
        "extraction_method": record.get("extraction_method"),
        "ocr_used": bool(record.get("ocr_used", False)),
        "classification_method": record.get("classification_method"),
        "matched_reference_sample": record.get("matched_reference_sample"),
    }

    supported_types = {"PURCHASE_ORDER", "SUPPLIER_INVOICE"}
    fields_to_check = ["supplier", "document_number", "document_date", "amount", "currency"]
    if document_type == "SUPPLIER_INVOICE":
        fields_to_check.append("po_reference")
    missing_fields = [field for field in fields_to_check if cleaned.get(field) is None]
    text_extracted = record.get("text_extracted", True)

    if not text_extracted:
        status = "UNREADABLE_DOCUMENT"
        note = "No usable text was extracted. Correct the document or verify OCR, then recheck."
    elif document_type == "UNKNOWN":
        status = "INTAKE_REVIEW"
        note = "The document type could not be determined reliably. A reviewer must classify it."
    elif document_type not in supported_types:
        status = "UNSUPPORTED_DOCUMENT"
        note = f"{document_type} is readable but outside the current PO and supplier-invoice scope."
    elif missing_fields:
        status = "READY_WITH_GAPS"
        note = "The evidence is readable and standardised. Module 2 must assess the missing information."
    else:
        status = "READY_FOR_MODULE_2"
        note = "The document was extracted and standardised successfully."
    if cleaned["ocr_used"] and status in {"READY_FOR_MODULE_2", "READY_WITH_GAPS"}:
        note = "OCR was used. " + note

    cleaned["processing_status"] = status
    cleaned["missing_fields"] = ", ".join(missing_fields)
    cleaned["intake_note"] = note
    return cleaned


def process_attachments(downloaded_files):
    """Extract, classify, standardise and filter content-equivalent duplicates."""
    create_folders()
    records = []
    config = load_reference_config()
    profiles = build_reference_profiles()
    content_registry = load_content_registry()

    for item in downloaded_files:
        file_path = Path(item["path"])
        email_id = item["email_id"]
        source_hash = item.get("source_hash") or calculate_hash(file_path.read_bytes())
        try:
            if file_path.suffix.lower() in {".csv", ".xlsx"}:
                raw_records, canonical_content = read_table(file_path, email_id, source_hash)
            else:
                text, extraction_method, ocr_used = extract_text_from_path(file_path)
                canonical_content = text
                raw_records = [
                    extract_fields_from_text(
                        file_path, email_id, source_hash, text, extraction_method,
                        ocr_used, config, profiles,
                    )
                ]

            fingerprint = calculate_content_fingerprint(canonical_content)
            existing = content_registry.get(fingerprint) if fingerprint else None
            if existing and existing.get("source_hash") != source_hash:
                log_duplicate(
                    email_id=email_id,
                    filename=file_path.name,
                    source_hash=source_hash,
                    duplicate_of=existing.get("source_file", "previously processed content"),
                    reason="NORMALIZED_CONTENT_MATCH",
                )
                continue

            cleaned_records = [clean_record(raw, fingerprint) for raw in raw_records]
            records.extend(cleaned_records)
            if fingerprint and fingerprint not in content_registry:
                register_content_fingerprint(fingerprint, source_hash, file_path.name, email_id)
                content_registry[fingerprint] = {
                    "source_hash": source_hash,
                    "source_file": file_path.name,
                }
        except Exception as error:
            records.append(
                {
                    "record_id": create_record_id(source_hash),
                    "document_type": "UNKNOWN",
                    "supplier": None,
                    "document_number": None,
                    "po_reference": None,
                    "document_date": None,
                    "amount": None,
                    "currency": None,
                    "source_file": file_path.name,
                    "source_extension": file_path.suffix.lower(),
                    "source_hash": source_hash,
                    "content_fingerprint": None,
                    "source_row": None,
                    "email_id": email_id,
                    "extraction_method": "PROCESSING_ERROR",
                    "ocr_used": False,
                    "classification_method": "PROCESSING_ERROR",
                    "matched_reference_sample": None,
                    "processing_status": "INTAKE_REVIEW",
                    "missing_fields": "document_type, supplier, document_number, document_date, amount, currency",
                    "intake_note": f"Processing error: {error}",
                }
            )
    return records


def ensure_standard_columns(dataframe):
    for column in STANDARD_COLUMNS:
        if column not in dataframe.columns:
            dataframe[column] = None
    return dataframe[STANDARD_COLUMNS]


def export_results(records):
    create_folders()
    new_data = ensure_standard_columns(pd.DataFrame(records, columns=STANDARD_COLUMNS))
    if INTAKE_FILE.exists():
        try:
            existing_data = pd.read_csv(INTAKE_FILE)
        except pd.errors.EmptyDataError:
            existing_data = pd.DataFrame(columns=STANDARD_COLUMNS)
    else:
        existing_data = pd.DataFrame(columns=STANDARD_COLUMNS)
    existing_data = ensure_standard_columns(existing_data)
    all_data = pd.concat([existing_data, new_data], ignore_index=True)
    all_data = all_data.drop_duplicates(subset=["record_id"], keep="last")
    all_data = ensure_standard_columns(all_data)

    module2_data = all_data[
        all_data["processing_status"].isin(["READY_FOR_MODULE_2", "READY_WITH_GAPS"])
    ]
    intake_review = all_data[
        all_data["processing_status"].isin(["INTAKE_REVIEW", "UNREADABLE_DOCUMENT"])
    ]
    unsupported = all_data[all_data["processing_status"] == "UNSUPPORTED_DOCUMENT"]

    all_data.to_csv(INTAKE_FILE, index=False)
    module2_data.to_csv(MODULE2_FILE, index=False)
    module2_data.to_json(MODULE2_JSON_FILE, orient="records", indent=2)
    intake_review.to_csv(REVIEW_FILE, index=False)
    unsupported.to_csv(UNSUPPORTED_FILE, index=False)

    print(f"Evidence received: {len(all_data)}")
    print(f"Complete extraction: {(all_data['processing_status'] == 'READY_FOR_MODULE_2').sum()}")
    print(f"Passed with gaps: {(all_data['processing_status'] == 'READY_WITH_GAPS').sum()}")
    print(f"Needs intake review: {len(intake_review)}")
    print(f"Unsupported documents: {len(unsupported)}")


def main():
    create_folders()
    print("Checking the email inbox...")
    downloaded_files = download_email_attachments()
    print(f"Unique downloaded attachments: {len(downloaded_files)}")
    if not downloaded_files:
        print("No new unique supported attachments were found.")
        return
    export_results(process_attachments(downloaded_files))


if __name__ == "__main__":
    main()
